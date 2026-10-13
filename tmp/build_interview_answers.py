# -*- coding: utf-8 -*-
"""Generate AI/ML Solution Architect interview answer document."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Cm

OUT = r"C:\Users\LENOVO\Downloads\Prasad_Nimmalapudi_AIML_Solution_Architect_Interview_Answers.docx"

NAVY = RGBColor(0x1B, 0x3A, 0x4B)
TEAL = RGBColor(0x0F, 0x6E, 0x6E)
BODY = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x4B, 0x55, 0x63)


def set_run_font(run, name="Calibri", size=11, bold=False, color=BODY, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def shade_cell(cell, hex_color):
    tc = cell._tePr if hasattr(cell, "_tePr") else cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def add_page_number(paragraph):
    run1 = paragraph.add_run("Page ")
    set_run_font(run1, size=9, color=MUTED)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run2 = paragraph.add_run()
    run2._r.append(fld_begin)
    run2._r.append(instr)
    run2._r.append(fld_end)
    set_run_font(run2, size=9, color=MUTED)
    run3 = paragraph.add_run("  |  Nimmalapudi Prasad  |  AI/ML Solution Architect")
    set_run_font(run3, size=9, color=MUTED)


def para(doc, text, size=11, bold=False, color=BODY, space_before=0, space_after=8, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color, italic=italic)
    return p


def mixed(doc, parts, space_before=0, space_after=8):
    """parts: list of (text, bold, italic, color, size)"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.12
    for text, bold, italic, color, size in parts:
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold, color=color, italic=italic)
    return p


def h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_run_font(r, size=16, bold=True, color=NAVY)
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "0F6E6E")
    pbdr.append(bottom)
    pPr.append(pbdr)
    return p


def question(doc, number, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    r1 = p.add_run(f"{number}. ")
    set_run_font(r1, size=12, bold=True, color=TEAL)
    r2 = p.add_run(text)
    set_run_font(r2, size=12, bold=True, color=NAVY)
    return p


def label_answer(doc, text):
    parts_iter = []
    # Split on **bold** markers
    buf = ""
    bold = False
    i = 0
    tokens = []
    while i < len(text):
        if text.startswith("**", i):
            if buf:
                tokens.append((buf, bold))
                buf = ""
            bold = not bold
            i += 2
        else:
            buf += text[i]
            i += 1
    if buf:
        tokens.append((buf, bold))
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.12
    first = True
    for chunk, is_bold in tokens:
        r = p.add_run(chunk)
        set_run_font(r, size=11, bold=is_bold, color=BODY)
        first = False
    return p


SECTIONS = []

# content tuples: ("q", n, question) or ("a", answer with **bold**) or ("h", heading) or ("note", text)

CONTENT = []

def add_h(t):
    CONTENT.append(("h", t))

def add_q(n, t):
    CONTENT.append(("q", n, t))

def add_a(t):
    CONTENT.append(("a", t))

add_h("How to use this document")
add_a("These are spoken answers for the AI/ML Solution Architect question bank. Each answer is written in the first person so you can practice it aloud. Keep design answers to about four minutes and screen answers to about ninety seconds. Draw the three pictures at the end while you talk: request path, data path, and release path.")
add_a("Use only the metrics already on your resume: about 15% developer productivity from an autonomous coding agent, about 40% less content-production time from a CrewAI and LangGraph multi-agent system, about 30% less database load from Redis and Memcached, program-spend forecasts within about 5%, about 20% shorter AI product cycle time, about 60% less sprint carryover, and the Email DataMart redesign from 7 hours 30 minutes to 2 hours 10 minutes. If a follow-up asks for a number you cannot defend, say what you measured and what you did not.")
add_a("React, TypeScript, Node.js, GraphQL, and Kafka are required by this job description and are lighter on the resume than Python, FastAPI, REST, Pub/Sub, and Kubernetes. Those answers are architect answers: you own the contract, the boundary, and the review bar. Do not claim you personally built a Ford React application. Pub/Sub is your messaging proof point. Describe Kafka as the same event-log pattern on the AWS or Azure estate this role uses.")

add_h("1. System design")

add_q(1, "Design an enterprise AI assistant that answers from internal policies and analytics, cites sources, enforces user permissions, and escalates to a human.")
add_a("I start from the user and the trust boundary, not from the model. A React and TypeScript client talks only to a Node.js backend-for-frontend. That BFF calls an API gateway. Behind the gateway, a FastAPI experience service authenticates the user, resolves entitlements, and opens a trace. The orchestrator is a LangGraph workflow with a hard step limit. It does four things: classify the question, retrieve, call a tool if needed, and either answer with citations or hand off to a human queue.")
add_a("Retrieval is hybrid search over Pinecone, filtered by the user’s groups before similarity runs. Postgres holds cases, entitlements, and workflow state. Redis holds the short session and a response cache. The system of record for policy text stays the enterprise content system. We index it. We do not copy it into the prompt store as the master. Analytics questions go through a governed semantic layer or a read API, not raw warehouse SQL from the model.")
add_a("The model is behind a gateway that can route to Azure OpenAI, Bedrock, Gemini, Claude, or Groq. Tools are MCP servers with scoped credentials. Every tool call and every citation is written to an audit topic. If retrieval confidence is low, or the question needs a decision the policy does not cover, the graph stops and creates a human case with the draft and the sources attached.")
add_a("I would deploy this on AKS or EKS with Helm, private networking, and a separate namespace for inference. At Ford I used this shape for agent workloads: MCP as the integration boundary, Docker and Kubernetes with Helm and autoscaling for MCP servers and inference, and Llama Guard so the model cannot expose data the user is not allowed to see. The outcome I would hold the design to is a grounded answer with a citation, under a latency SLO, with a refusal when the evidence is weak.")

add_q(2, "Split that system into microservices. Which services exist on day one, and what stays a module?")
add_a("I split by data ownership, scaling, and failure, not by framework. On day one I want six deployables. The BFF owns the UI contract and aggregation. The experience API owns auth context and case creation. The orchestrator owns the graph, prompts, and step limits. The retrieval service owns chunking, embeddings, and Pinecone. The model gateway owns provider routing, budgets, and fallback. Each MCP server is its own deployable because each enterprise system has its own credentials and release cycle.")
add_a("I keep prompt templates, evaluators, and the abstain rule inside the orchestrator until they need an independent release. I do not create a separate “prompt microservice” in month one. The audit writer can start as a library that publishes to Kafka or Pub/Sub, and become a service when volume or retention rules demand it.")
add_a("The test I use: if this part can fail, scale, or be released without the others, it is a service. If it only changes when the workflow changes, it is a module. That is how I kept agent platforms releasable at Ford without a mesh of tiny services nobody could operate.")

add_q(3, "A legacy analytics platform, two internal REST APIs, a GraphQL partner API, and an external LLM must ship as one product. Show the integration map.")
add_a("I put an anti-corruption layer in front of each system. The analytics platform gets a read adapter that exposes approved metrics, not tables. Each internal REST API gets a client with timeouts, retries, and a circuit breaker. The partner GraphQL API is wrapped by the Node.js BFF or a dedicated adapter so the agent never sees the partner schema. The external LLM is reached only through the model gateway.")
add_a("Synchronous calls are for the user request path, and only when the answer is blocked on that data. Everything else is async. A canonical event, such as CaseUpdated or DocumentIndexed, is published to Kafka or Pub/Sub. Consumers index, notify, and audit. The canonical model is ours. Partner and legacy field names stop at the adapter.")
add_a("I have done the two hard parts of this already. At Ford I standardized external tools and APIs behind MCP so agents did not grow one-off clients. I also put Redis and Memcached in front of hot reads and cut database load by about 30%. For the analytics side, my IBM work on BigQuery, Dataflow, Dataproc, and Pub/Sub is the pattern I reuse: pipelines feed the warehouse, and the assistant reads a governed mart, not the raw feed.")
add_a("Operational rules: idempotency keys on any write, a schema registry for events, and a deadline budget so the LLM call cannot consume the whole latency SLO. If the partner GraphQL API is down, the product degrades to policy answers and a clear gap, rather than failing the whole request.")

add_q(4, "Add a second region. What is stateful, what is replicated, and what happens if Kafka or the vector index lags?")
add_a("I separate three kinds of state. Request state in Redis is regional and disposable. If the region dies, the user retries. Postgres for cases and entitlements is the system of record. I run it as a primary in one region with a synchronous or near-synchronous replica, and I promote with a documented runbook. Pinecone or the vector index is a derived store. I can rebuild it from the document source, so I replicate for read latency, and I accept a lag measured in minutes as long as answers cite the version they used. Kafka or Pub/Sub is the replay log. Consumers are idempotent, so a lag delays indexing and audit. It does not corrupt the case record.")
add_a("The recovery point I promise: no loss of committed cases. The recovery time I promise: a regional failover of the stateless tier in minutes, and a database promotion on a rehearsed path. I do not promise zero lag on the vector index. I show the index version on the citation so a stale chunk is visible. Cost follows the same split. GPU or high-rate inference stays in the primary region until the traffic justifies a second pool.")

add_q(5, "The same platform must run on Azure and AWS. What do you standardize, and which services stay native?")
add_a("I standardize the workload, not the cloud console. Containers, Helm charts, Terraform modules, the Kafka or Pub/Sub client, object storage, and OIDC are the portable core. GitHub Actions or GitLab CI builds one image. ArgoCD or Flux deploys it to AKS and to EKS. The model gateway hides Azure OpenAI versus Bedrock versus a self-hosted route.")
add_a("I leave these native: AKS versus EKS, Azure API Management versus AWS API Gateway, Key Vault versus Secrets Manager, Azure Monitor versus CloudWatch, and the private endpoint pattern on each network. Abstracting those creates a lowest-common-denominator platform that operators cannot debug.")
add_a("My cloud work spans Azure, AWS, and GCP, including BigQuery, Dataflow, and Pub/Sub on the data side, and EKS, AKS, and GKE on the runtime side. In an interview I pick one primary landing zone for this role and keep the other as a disaster-recovery or customer-residency target. I say which identity system is the source of truth, because dual identity is where these programs stall.")

add_h("2. AI/ML and agent architecture")

add_q(6, "A stakeholder says “use GPT for everything.” When do you choose classical ML, RAG, fine-tuning, or an agent?")
add_a("I choose from the job, not from the model brand. If the question is a stable prediction on structured features, I use classical ML. At Ford, program-spend forecasting within about 5% was that kind of problem: features, a model, monitoring for drift, and a batch or online score. A chat model is the wrong tool there.")
add_a("If the answer must come from documents that change, I use RAG. Policy, procedures, and citations belong here. If the format or tone is stable and prompting plus retrieval still fails an eval set, I consider fine-tuning or a small adapter, and only after the eval proves the gap. Fine-tuning is not how I keep policies current.")
add_a("If the work is multi-step, crosses systems, and branches, I use an agent. That is the CrewAI and LangGraph content system that cut production time by about 40%, and the autonomous coding agent that improved developer productivity by about 15%. The agent gets a tool allow-list, a step cap, and a stop condition. GPT, Claude, Gemini, or Groq is the reasoning engine inside one of these four choices. It is not the architecture.")

add_q(7, "Design the RAG pipeline from ingest through citation and abstain. Where do access labels get enforced?")
add_a("Ingest starts in the source system, with the document’s access labels attached at extraction time. I parse by structure, heading and section, not by a blind character count. Chunks inherit the document ACL, the source URI, and a version. Embeddings are written to Pinecone with those fields as metadata. The embedding model is versioned with the index. A model change is a reindex, not a silent overwrite.")
add_a("At query time, the user’s groups are resolved in the experience API. The Pinecone filter applies those groups before vector similarity. A later prompt instruction cannot widen that filter. I use hybrid search, then a reranker, then a context budget. The generator must cite chunk IDs. If the top score is below a threshold, or the citations do not support the claim, the orchestrator abstains and offers a human handoff.")
add_a("Access control is enforced in retrieval, again in the tool layer, and checked on the way out. Llama Guard, or the same class of policy, screens input and output for injection and for data the user should not see. I learned to treat this as a product control at Ford, not as a filter we turn on after the demo.")

add_q(8, "How do you evaluate the assistant before production?")
add_a("I refuse to promote on a demo. I keep a golden set of questions with expected sources, expected refusals, and expected tool calls. Retrieval is scored with precision and recall against labeled chunks. Generation is scored for groundedness: every claim maps to a citation. I also track refusal quality, so we do not celebrate a model that abstains on easy questions. Tool use is scored as exact match on the allow-listed call, including arguments.")
add_a("Offline eval runs in CI when prompts, tools, chunking, or the model route change. A drop beyond an agreed delta blocks the release. Online, I sample production traces, score them on a delay, and watch citation-open rate and human-override rate. At Ford I used precision, recall, and F1 on model evaluation, plus prompt iteration. I would add a groundedness set for this assistant because accuracy on a classifier is not the same as “the policy sentence is actually in the source.”")

add_q(9, "Compare a ReAct loop, LangGraph, and CrewAI. Where do you cap steps and require approval?")
add_a("ReAct is a short loop: think, call one tool, observe, answer. I use it when the path is one or two lookups and a draft. It is the wrong default for a workflow that must pause for a person.")
add_a("LangGraph is what I want in production when state, retries, and human approval matter. Nodes are explicit. The graph persists state in Postgres. A node can be “wait for approver.” I can resume after a failure without replaying side effects. I cap steps, cap tool calls, and cap tokens. Crossing the cap routes to a human with the trace attached.")
add_a("CrewAI fits when the work is genuinely role-shaped. Researcher, writer, and reviewer passing a packet is the pattern I used to cut content production time by about 40%. I still put a supervisor limit on it. Roles do not get an open tool list. Approval is required before any tool that writes, sends, or spends money. Read-only retrieval can proceed. A refund, an email to a customer, or a policy publish cannot.")

add_q(10, "Where does MCP sit? Design one MCP server for an enterprise system.")
add_a("MCP is the integration boundary between the agent and the enterprise. The agent does not get a database password or a hand-written client per system. It calls tools on an MCP server that I own and release.")
add_a("Take an HR policy system as the example. One MCP server, one deployable, its own service account. Tools are few and schema-validated: search_policy, get_policy_section, create_case. search_policy requires the user token, and the server applies entitlements on the query. get_policy_section returns text plus a citation handle, not a file dump. create_case is a write, so the graph must have passed an approval node, and the call carries an idempotency key. Every call is audited with user, tool, arguments hash, and result code. Credentials live in Key Vault or Secrets Manager and are mounted at runtime. The server is stateless and scales with Helm on AKS or EKS.")
add_a("At Ford I standardized external tool and API integrations with MCP so new agents reused the same servers. We deployed those servers and the inference services on Docker and Kubernetes with Helm and autoscaling. That is the design I would repeat here, with a tighter tool allow-list and user-scoped auth on every call.")

add_q(11, "You have OpenAI, Claude, Gemini, and Groq. Design the model gateway.")
add_a("Application code calls one internal API: complete, embed, and moderate. It never imports a vendor SDK. The gateway holds the route table, the prompt and model version, token and cost budgets per tenant, timeouts, and fallback. A policy might send a short classification to Groq for latency, the final grounded answer to Claude or Azure OpenAI, and embeddings to a single pinned model so the Pinecone index stays compatible. Gemini is a fallback region or a second opinion, not a silent switch in the same index.")
add_a("Fallback is for transport errors and rate limits, not for a low-quality answer. Quality fallback is an eval decision, made offline. Every call logs provider, model version, tokens, latency, and cost onto the trace. A budget alarm trips before the invoice does. I have operated this idea as LLMOps across OpenAI-class models, Claude, and the gateway pattern. Groq is the low-latency lane. The stronger model is the answer lane. The orchestrator does not know which one ran, except in the audit record.")

add_q(12, "An agent hallucinates a policy and calls the wrong API. What guardrails run?")
add_a("I treat this as three gates. On input, I screen for prompt injection and for attempts to override the tool policy. Llama Guard, or an equivalent classifier plus a deny-list, runs before the graph spends a tool call. In the middle, the orchestrator can call only registered MCP tools. Arguments are schema-checked. A write tool is absent from the graph unless an approval node has passed. Retrieval filters are applied in the index, so a persuasive prompt cannot fetch another user’s policy. On output, I require citations that resolve to retrieved chunk IDs. If a sentence has no citation, it is dropped or the whole answer abstains. A second policy check looks for secrets and for cross-tenant strings.")
add_a("At Ford I put Llama Guard-style guardrails in place specifically to reduce hallucination and unauthorized data exposure. The operating rule I would state in the interview: a fluent answer with no source is a defect, and a tool call that is not on the allow-list is an incident, not a retry.")

add_q(13, "Design offline and online features for a forecast model that must stay within about 5% error, next to the GenAI stack.")
add_a("The forecast and the assistant share definitions. They do not share a database connection. Feature computation for program spend sits in the warehouse path I know from GCP and from enterprise ETL: Dataflow or Airflow builds the training table, the model trains on a schedule, and batch scores land in a mart. Online, a feature service or a cached mart read serves the latest score to the application. Drift, data freshness, and error against actuals are the monitors. The 5% band is an SLO on the forecast, reviewed with finance, not a claim the chatbot repeats.")
add_a("The agent may read the published score through an MCP tool or a read API. It may explain the drivers if those drivers were stored. It may not run ad hoc SQL over the finance warehouse, and it may not retrain the model. At Ford, AI-driven forecasting stayed inside about 5% and allowed budget adjustments in time. I would describe that model as a product with an owner, a training set, and a rollback to the last good score, sitting beside RAG rather than inside it.")

add_h("3. APIs, data, and events")

add_q(14, "REST versus GraphQL. What does the React client call, and what is forbidden from the browser?")
add_a("The browser calls one origin: the Node.js BFF. It does not call MCP, Pinecone, the model gateway, or the warehouse. Where the UI needs a stable command, the BFF exposes REST: create case, submit feedback, approve a draft. Where the UI must assemble a case, the user’s entitlements, and the citations in one round trip, the BFF exposes a small GraphQL schema it owns. Resolvers call internal REST services. I do not let the partner’s GraphQL schema leak into our UI.")
add_a("Auth is at the gateway and again in each service. The GraphQL layer is not a place to hide a missing entitlement check. I watch for N+1 queries and I cap query depth and cost. My deepest implementation experience is Python and FastAPI on the service side, plus API gateway and REST integration. For this role I own the BFF contract, the typed client the React team consumes, and the review rule that no token with tool scope ever ships in the frontend bundle.")

add_q(15, "Design the Kafka topics for document indexed, agent run completed, tool failed, and audit.")
add_a("I use four topics, keyed by tenant and then by aggregate id, so one tenant’s documents stay ordered. document.indexed is emitted by the ingest worker after Pinecone acknowledges the version. Consumers refresh caches and notify the UI. agent.run.completed carries the run id, outcome, citation ids, token cost, and latency. It does not carry the full prompt. tool.failed carries the tool name, error class, and whether the call was retried. audit.events is append-only, longer retention, and restricted ACLs. Payloads with document text go to object storage. The event holds a pointer.")
add_a("Consumers are idempotent, using the event id. Retries are bounded. A dead-letter topic catches poison messages, and a person owns that queue. I use an outbox from Postgres for case events so a committed case cannot miss its event. My production messaging proof point is Pub/Sub in the GCP data pipelines at IBM, with the same rules: schema, retry, and replay. On this role I implement that design on Kafka because the estate is Azure and AWS and the consumers need a durable log, not only a push queue.")

add_q(16, "SQL and NoSQL for this platform. What is the system of record?")
add_a("Postgres is the system of record for users, cases, approvals, and workflow state. I want transactions and constraints there. Redis is the cache and the short-lived agent session. It is not a record. Pinecone is the vector index, derived from source documents, rebuildable. A document database can hold large run blobs if Postgres row size becomes a problem, with the case id still mastered in Postgres. SQLite or UnQLite is for local developer or edge agent state on a laptop. I do not put them in the production request path.")
add_a("The source policy system remains the record for policy text. Our index is a copy with a version. At Ford, Redis and Memcached in front of the database cut load by about 30%. That is the cache lesson I keep: cache the hot read, name the TTL, and invalidate from the write path or the index-completed event. I do not cache entitlements for long.")

add_q(17, "A partner API times out and a consumer reprocesses the same event. How do you keep side effects from running twice?")
add_a("Every write carries an idempotency key derived from the business action, not from the delivery attempt. create_case uses tenant plus client request id. A tool call uses run id plus tool name plus a hash of the arguments. The MCP server stores the key and the result. A retry returns the stored result. The consumer treats “already applied” as success.")
add_a("Timeouts get a special rule. If I do not know whether the partner committed, I do not blindly retry a non-idempotent call. I reconcile with a read, or I retry only if the partner honors the same key. The Kafka consumer commits the offset after the handler acknowledges, and the handler itself is safe to run twice. At the user level I would rather show “case already created” than create two cases. This is the same discipline I apply to agent tools: a write without an idempotency key is not allowed on the allow-list.")

add_q(18, "A GraphQL field and a Kafka event both change in one release. How do you version them?")
add_a("I add fields. I do not break old consumers in the same release. GraphQL gets a new nullable field. The old field stays until clients have moved, then a deprecation window ends it. Kafka gets a new schema version in the registry, backward compatible. New consumers understand both versions for one release. The producer ships first only if the new schema is optional. If I must break compatibility, it is a new topic or a new event type, and both run until the lagging consumer is gone.")
add_a("The release train is one Helm release for the producer and a feature flag so I can stop emitting the new shape. Contract tests in CI fail the build if the BFF schema or the event schema removes a field a known client still reads. I treat prompt and tool schema changes the same way: they are versioned contracts, not strings we edit in production.")

add_h("4. Cloud, Kubernetes, and delivery")

add_q(19, "Walk a service from commit to production. Where do tests and evals gate the release?")
add_a("A commit to the service repo runs GitHub Actions or GitLab CI. Unit tests, contract tests, and a container build come first. The image is scanned and signed. Helm packages the chart. On the agent and retrieval services, an offline eval set runs against the candidate prompt, tool schema, or model route. A groundedness or retrieval regression beyond the delta fails the pipeline. ArgoCD or Flux syncs the chart to dev, then test. Production is a promotion of the same digest, not a rebuild. A canary takes a slice of traffic. The SLO and the cost burn are watched. Then we step up.")
add_a("At Ford I cut AI product cycle time by about 20% by making this path the way work shipped, Agile planning included, and I cut sprint carryover by about 60% when code generation and automated root-cause were inside the delivery loop rather than beside it. The architectural point for this interview: a prompt change is a release. It does not skip the eval gate because no Python file changed.")

add_q(20, "Kubernetes for MCP servers and LLM inference. How do you size, scale, and contain cost?")
add_a("MCP servers are normal stateless services: requests and limits set from a load test, HPA on CPU and on a custom metric such as in-flight tool calls, a pod disruption budget, and a network policy that allows traffic only from the orchestrator. Secrets come from Key Vault or Secrets Manager through a driver, not from the image.")
add_a("Inference is a different pool. If GPUs are in the cluster, they sit on a dedicated node pool with taints. If the model is Azure OpenAI or Bedrock, the “inference” deployment is the gateway, and the scale signal is queue depth and token rate, not GPU utilization. HPA scales the gateway and the MCP tier. VPA is only for steady workloads where I am right-sizing, not for spiky agent traffic. A namespace quota and a monthly cost budget alarm are part of the chart’s operating contract. I have run this on EKS, AKS, and GKE, with Helm, Kustomize, and autoscaling, including the Ford agent workloads.")

add_q(21, "Terraform the landing zone. What is an account guardrail versus an app default?")
add_a("The landing zone is account or subscription level, and applications cannot opt out. I put the network (hub, private endpoints, no public cluster API if the standard says so), the identity baseline (OIDC for CI, workload identity for pods), the log archive, the key vault, and policy that blocks public buckets and open security groups. Terraform modules for those live in a platform repo. App teams consume them. They do not fork them.")
add_a("App defaults are the Helm values: replica count, CPU, the model route, the topic names. Those can vary by environment. Guardrails cannot. I have used Terraform, Ansible, and GitOps with ArgoCD and Flux. My rule in a review is simple. If a mistake affects every team or creates a public data path, it is a landing-zone control. If it affects one service’s latency, it is a chart value.")

add_q(22, "Where do you terminate auth, rate limits, mTLS, and canary traffic if Istio or an API gateway is present?")
add_a("The external API gateway, Azure API Management or AWS API Gateway, terminates user auth, issues the rate limit the product sold, and validates the public schema. Istio, or the mesh, handles east-west mTLS, service identity, and the canary split between orchestrator versions. I do not do user login inside the mesh, and I do not rely on the mesh as the only quota on a public API.")
add_a("Ingress or NGINX is the edge only if the platform standard is ingress rather than a managed gateway. One of them owns TLS for the public name. Internal services still use mesh identity. Canary is a weight on the orchestrator deployment, with a header override so I can test a prompt release without exposing all users. Tool servers stay off the public ingress entirely.")

add_q(23, "A Helm release of a new prompt and tool schema is bad. How do you roll back, and what do you refuse to roll back?")
add_a("The release is the chart plus the config version that pins the prompt and the tool schema. ArgoCD rolls the deployment back to the previous digest in one action. Feature flags let me disable the new tool without a second deploy if the binary is fine and the schema is not. In-flight graphs that already took a write are not reversed by a rollback. I reconcile those runs from the audit log: complete, compensate, or hand to a person. I do not roll back Postgres case data or the audit topic. I do not roll back a Pinecone index by deleting vectors in a panic. I pin queries to the previous index version if I built the index side by side.")
add_a("The practice underneath this is the same one I use for SLO work. A bad release is an incident with a timeline, not a quiet kubectl edit. The next day the eval set gains the case that escaped.")

add_h("5. Observability and incident response")

add_q(24, "Define the observability stack. What is an SLO, and what pages a human?")
add_a("I want three layers, and I want them correlated by a run id. Metrics in Prometheus, dashboards in Grafana: latency, error rate, saturation, token rate, retrieval empty-rate, tool-error rate. Logs in ELK or OpenSearch, structured, with the run id and no raw document text. Traces in Datadog or the OpenTelemetry backend: one trace per user request. Alerts fire on burn rate against an SLO, not on every error log.")
add_a("A service SLO I would propose: p95 end-to-end under a budget we set with the product, excluding the human queue, and availability of the experience API. A product SLO: share of answers that are grounded and cited, measured on a sample. A cost SLO: tokens per successful task under a ceiling. A page goes out when the fast burn rate says the monthly error budget will be gone in hours, or when tool-failure or cost rate breaks a hard cap. A slow miss goes to a ticket. I have run this with Prometheus, Grafana, Datadog, ELK, SLO and SLI, and error budgets, including on-call. The agent does not get a pass on operability because the failure mode is a wrong sentence instead of an HTTP 500.")

add_q(25, "Cost and error rate spike after a model-routing change. What do you do?")
add_a("In the first fifteen minutes I stop the bleeding. I shift the canary to zero or flip the route back to the previous model in the gateway. I cap tokens per request if a loop is the cause. I open the dashboard for token rate, error rate, and the new route’s traffic share. I do not start a redesign on the call.")
add_a("In the first hour I read traces for a sample of failures. I look for a retry storm, a tool schema mismatch, or a fallback that called the expensive model on every request. I confirm the rollback is the previous known-good config, I check the error budget, and I write the timeline: change, detection, mitigation. The next day I add the failing case to the eval set, and I add a gateway alarm on cost per minute that would have caught this before the invoice. I have used on-call and incident response as part of the SRE practice around these platforms. The stakeholder update is one paragraph: impact, what users saw, what we reversed, what we will prevent.")

add_q(26, "Which spans must exist on one user request?")
add_a("One trace, parent span at the BFF or gateway. Child spans: experience API, orchestrator node for each graph step, retrieval with filter and hit count, model gateway with provider, model version, and token counts, each MCP tool with name and status, and database calls that dominate latency. I put the run id, tenant, prompt version, index version, and citation count on the trace as attributes. I do not put the document body or the full prompt if it contains employee data. Those go to a restricted store with a shorter retention, linked by id.")
add_a("That set is what lets me explain a wrong answer the next day. I can see whether retrieval returned the right section, whether the model ignored it, or whether the tool call never happened. Without those spans, the incident review becomes an argument.")

add_q(27, "The error budget is exhausted mid-sprint. What ships, and what waits?")
add_a("Reliability work ships. Feature work that touches the request path waits. I say that in the planning meeting with the budget chart, not as a surprise at the end of the sprint. Allowed changes are the rollback hardening, the bad-route fix, the eval gap, and a capacity change if saturation caused the burn. A new tool or a new prompt waits until the burn rate is back inside the policy for a week, unless it is the fix.")
add_a("I have managed this as error-budget policy, not as a mood. At Ford, protecting completion and cutting carryover mattered, and the way to protect a sprint is to stop adding risk when the system is already spending the budget the users were promised. Product hears a date for the feature, tied to the budget recovering, and they hear what I will not trade: the groundedness and availability SLOs.")

add_h("6. Security, governance, and compliance")

add_q(28, "Threat-model the assistant: prompt injection, cross-user retrieval, secret leakage, and over-broad tools.")
add_a("Prompt injection is untrusted text in a document or a user message that tries to rewrite the tool policy. The control is a fixed system instruction, a tool allow-list outside the model, and an input screen. The model does not grant itself permissions. Cross-user retrieval is a missing filter. The control is ACL metadata on every chunk and a mandatory filter from the user token before search. I test it with a case that must return zero hits. Secret leakage is a document or a prompt that contains a key, or a tool result that returns one. The control is secret scanning on ingest, redaction, output filtering, and a ban on logging raw payloads. Over-broad tools are an admin API exposed as one MCP method. The control is small tools, scoped service accounts, schema validation, and approval before writes.")
add_a("I would put these four in the architecture review checklist. At Ford, Llama Guard and the MCP boundary were how I kept agents from wandering into data and actions they were not given. A control that exists only in the prompt is not a control.")

add_q(29, "API management standards for REST, GraphQL, and MCP.")
add_a("Every API has an owner, a schema, and an auth scheme. External and UI-facing REST and GraphQL sit on the API gateway with OAuth2, scopes per operation, quotas per client, and request validation. Breaking changes follow the versioning rule from the schema question. Internal service calls use workload identity and mTLS. MCP is not published on the public gateway. It is an internal API with its own scope per tool, its own quota, and an audit record of every call. GraphQL gets depth and complexity limits so one query cannot become a denial of wallet or a data dump.")
add_a("Compliance is the retention schedule on audit, the data-processing terms for each model provider, and the list of fields we are allowed to send to a model. I write that list down. “We use enterprise mode” is not the standard. The standard is which fields leave the network, for how long, and under which contract.")

add_q(30, "A business unit wants tenant documents sent to a public model API. What is allowed?")
add_a("Public, no-contract endpoints are not allowed for tenant documents. I will allow a provider only where we have a contract that states no training on our prompts, a region we accept, and a deletion story. On Azure that is typically Azure OpenAI in the tenant region with private networking. On AWS it is Bedrock with private connectivity and the same terms. If the document is highly sensitive, the answer is a private deployment or a redacted excerpt, not the full file.")
add_a("What I offer the business unit is a path that still ships: approved provider, private endpoint, retrieval of the minimum chunks, and an audit of what was sent. What I will not offer is a developer key in a notebook against a consumer API. I have had to make this distinction in enterprise AI programs. The architect’s job is to show the compliant design that meets the outcome, and to be specific about the clause that blocks the shortcut.")

add_q(31, "Write the release checklist for a new agent tool.")
add_a("The tool has an owner and an on-call. The schema is reviewed: inputs, outputs, errors, and idempotency for writes. AuthZ is user-scoped, tested with an allow case and a deny case. The service account cannot do more than the tool says. An eval delta shows the new tool improves task success and does not increase uncited answers or wrong-tool calls. Cost per run is estimated and a budget alarm exists. The audit event is present. The Helm release has a rollback. A runbook says how to disable the tool with a flag. Documentation for the calling team is a short ADR plus a sample trace. I do not approve the tool because the demo looked good. I approve it because those boxes are checked.")

add_h("7. Leadership and the architecture story")

add_q(32, "Give a two-minute architecture narrative of a system you owned.")
add_a("At Ford Technology Services I was the AI Architect for production agent and platform work. The business problem was that expert workflows, content production, engineering support, and program planning were slow, manual, and inconsistent, while any automation had to be safe enough for an enterprise.")
add_a("The constraints were familiar. Multiple systems of record, no appetite for a model with open credentials, and a need to run on Kubernetes with a real release process. I set the architecture as a workflow, not a chatbot. LangGraph or CrewAI for the multi-step work, MCP servers as the only door into enterprise tools, a model behind a controlled runtime, and Docker, Kubernetes, and Helm with autoscaling so the team could operate it. Redis and Memcached took hot reads off the database. Evaluation used precision, recall, and F1, and Llama Guard reduced unsafe or ungrounded output. Delivery was Agile, with GenAI used inside the engineering loop for code generation and root-cause, not only in the product.")
add_a("The tradeoff I would name: we standardized on MCP and an explicit graph, which was slower to start than a single prompt, and much safer to extend. The results I will stand behind are about 40% less content-production time, about 30% less database load, about 15% developer productivity on the coding agent, forecasts within about 5%, about 20% shorter delivery cycle, and about 60% less sprint carryover. If they want one diagram, I draw the orchestrator, the MCP boundary, the cache, and the cluster.")

add_q(33, "You standardized how teams integrate tools for agents. Who disagreed, and what improved?")
add_a("The disagreement is usually speed. A team with one API wants a direct client inside the agent so they can demo on Friday. MCP looks like a tax. I agreed that a direct client is fine for a spike on non-production data. I did not agree that the spike becomes the production path. The production rule was one MCP server per system, schema-validated tools, and shared credentials management.")
add_a("What I traded was a week of extra work on the first server, in exchange for the second and third agents not rebuilding auth and retries. What improved was reuse and reviewability. A security review looked at the server once. New use cases were new tools or new graphs, not new integrations. That is the Ford standardization I would tell as a stakeholder story: early friction, then a platform other teams could adopt without a meeting for every API.")

add_q(34, "Mentor a senior engineer into an architect. What do you put in their hands in the first month?")
add_a("I give them one vertical slice they own end to end, not a tour of every system. In the first month that is a reference MCP server, the Helm chart that deploys it, the eval harness that gates it, and an architecture decision record for one real tradeoff they had to make. I review the ADR with them before they publish it. I sit in one design review as their second, then they lead the next one. Knowledge transfer is that repo and a one-hour walkthrough recorded for the team, plus office hours. Slides are optional. The standard I am teaching is the one I want them to enforce: contracts, operability, and a metric, not a framework preference.")

add_q(35, "Product wants a chatbot in four weeks and wants to skip evaluation and access control. What do you ship?")
add_a("I ship a thin slice on non-sensitive, already-public or broadly readable content, with citations, logging, and a small eval set that must pass. I do not ship tenant-confidential documents, and I do not ship write tools. Access control is not a later hardening story. If the index cannot filter by user, those documents stay out. The four-week outcome is a cited assistant on the safe corpus, a human handoff, and a date for the confidential corpus once the filter and the eval are in place.")
add_a("I say this as a choice that still delivers a demo the business can use, rather than a refusal to start. I have had to hold this line because a fluent wrong policy is worse than a late feature. The sentence I practice: we can show value in four weeks on content we are allowed to show, and we will not skip the control that keeps one employee from seeing another employee’s data.")

add_q(36, "Two teams want different clouds and different agent frameworks. How do you set the standard?")
add_a("I set the seams, not the slogan. The standard is the model gateway contract, the MCP tool contract, the trace and audit fields, the Helm runtime, and the eval gate. Inside that, a team can use LangGraph or CrewAI if they can meet the step cap, the allow-list, and the on-call. A second cloud is allowed when there is a residency or acquisition reason, and the landing-zone guardrails still apply. A second cloud because of preference is a platform tax, and I will say so with the cost of two identity models and two clusters.")
add_a("I influence this by a reference implementation, not by a mandate email. The team that wants a different framework adopts the seams on one workflow first. If they cannot meet the operability bar, they use the reference stack. That is the same stakeholder pattern I used to align engineering, product, and architecture around a GenAI path at Ford.")

add_q(37, "Tell a stakeholder story in business terms, with no framework names until they ask.")
add_a("I would use the content-production result. Our teams were spending too long producing material that followed the same expert steps every time. I led an approach where software drafted the work, checked it against our sources, and left the judgment calls to people. We cut the production time by about 40%. The way we got there was a controlled workflow with a stop when the source was missing, not an open-ended chat. When they ask how, I then say CrewAI and LangGraph, MCP for the systems we had to touch, and a review step before anything was published.")
add_a("If the room is engineering leadership, I switch to the database story: hot reads were loading the system of record, we introduced a distributed cache, and database load dropped by about 30%, which protected the applications in front of it. If the room is finance, I use the forecast that stayed within about 5% and let them adjust spend in time. I pick one story. I do not recite the resume.")

add_h("8. Screen questions")

add_q(38, "Why Python at the agent and model boundary, and Node.js at the BFF?")
add_a("Python is where the model, retrieval, and data libraries are strongest. FastAPI is the service I use for the orchestrator, retrieval, and MCP-adjacent APIs because the same engineers own the evals, the dataframe work, and the model gateway. Node.js is the BFF because the React and TypeScript client wants a typed, low-friction aggregation layer, and GraphQL resolvers are a natural fit there. I keep business rules and entitlements in the Python services. The BFF aggregates and shapes. It does not become a second orchestrator. My hands-on depth is Python and FastAPI. I can design and review the Node boundary and I staff the BFF with people who write Node every day.")

add_q(39, "What have you personally built in React and TypeScript, and what did you own versus review?")
add_a("I answer this without stretching Ford agent work into a frontend claim. My ownership is the API the UI must consume: the REST and GraphQL contract, streaming behavior for tokens, how citations are structured, auth expiry, and the rule that keys and tool scopes never ship to the browser. I review the React and TypeScript implementation against that contract: loading and error states, no direct calls to the model or the index, and accessible rendering of sources. Where I have written UI, I describe that work at the level I actually did it. I do not invent a production React codebase. The panel is hiring an architect who can lead this estate. They will trust a precise boundary more than a borrowed story.")

add_q(40, "One Azure service and one AWS service for private LLM access, and why.")
add_a("On Azure I choose Azure OpenAI behind a private endpoint, with Key Vault for keys and API Management in front. The reason is network isolation plus the enterprise contract on data use, inside the tenant we already govern. On AWS I choose Amazon Bedrock with private connectivity and IAM roles for the workload, secrets in Secrets Manager. The reason is the same: the application role calls the model, there is no long-lived key in the image, and the traffic stays on the private network. In both cases the model gateway is ours, so the product can fail over or switch providers without a UI change. I name GCP, Vertex, and Pub/Sub only if they ask about the broader multi-cloud work, so the answer stays on the two clouds this job cares about.")

add_q(41, "Docker Compose on a laptop versus Kubernetes in production. What changes?")
add_a("Compose is how I run the orchestrator, a local MCP server, Postgres, and Redis on a laptop so an engineer can demo a path. It is not the production control plane. In production I need rolling deploys, health checks, horizontal autoscaling, secrets from a vault, network policy, a disruption budget, and integration with the GitOps pipeline. Helm carries the values that Compose fakes with a local file. Identity changes too: a laptop uses a developer credential, a pod uses workload identity. I have used Docker Compose for local stacks and Kubernetes, Helm, and Kustomize for the environments that page me. The mistake I call out in reviews is a Compose file promoted to production because it “already works.”")

add_q(42, "One SQL performance problem you fixed, and one cache you introduced.")
add_a("The cache story is the crisp one. Hot read paths were hammering the system of record. I introduced distributed caching with Redis and Memcached, with a clear TTL and invalidation, and database load dropped by about 30%. I would explain the choice: cache what is read often and changes on a known write, and do not cache authorization decisions for long. On SQL itself, the Email DataMart redesign is my processing-time story: a solution redesign cut the run from 7 hours 30 minutes to 2 hours 10 minutes. If they want a query-level detail, I stay with what I personally changed, indexes, partitioning, or the job design, and I do not invent a query plan I cannot draw. Earlier, DataStage and Hadoop connectivity work at IBM was the same instinct: measure the slow stage, then fix the pipeline rather than the symptom.")

add_q(43, "Kafka versus a queue such as SQS or Service Bus. When is the log the right choice?")
add_a("A queue is right when one consumer group must do a job and then forget the message: send an email, kick a worker, retry a task. SQS or Service Bus is simpler there. A log is right when several consumers need the same fact, when I may add a consumer later, or when I must replay. Document indexing, audit, and “rebuild the vector index” are log problems. Kafka gives me retention, ordering by key, and independent consumer groups. My lived messaging system in the data platforms was Pub/Sub, which I used for pipeline events. I map that experience onto Kafka for this role: topics, keys, idempotent consumers, and a dead-letter path. I choose the queue when there is one worker and no replay story, so I do not operate a cluster I do not need.")

add_q(44, "How do you explain RAG to a non-technical director in four sentences?")
add_a("The assistant does not memorize the policy manual. When someone asks a question, we first find the few sections that person is allowed to see. We then ask the model to answer using only those sections and to show them. If we cannot find a solid source, it says so and offers a person, rather than guessing. That is why a policy change is a publishing change, not a months-long retraining project.")

add_q(45, "What do you refuse to put in an agent’s tool allow-list?")
add_a("I refuse any tool that is unbounded, irreversible, or wider than the user. That means raw SQL against a system of record, a shell on a production host, a “delete” or “send” with no idempotency key and no approval node, an admin API that can act as any user, and a tool whose only description is “call this URL with a free-form body.” I also refuse a tool that returns entire documents when a section would do. Reads of approved, filtered data can be on the list. Writes wait for a named approver in the graph. If a team cannot explain the worst thing the tool can do, it does not ship. That rule is how I keep an agent inside the architecture instead of letting it become an unauthenticated integration layer.")

add_h("Whiteboard pictures to draw while you talk")

add_a("**Request path.** React and TypeScript client, Node.js BFF, API gateway, FastAPI experience service, LangGraph orchestrator, retrieval, model gateway, MCP tools, audit. Say which box is public and which box is private.")
add_a("**Data path.** Source systems, ingest with access labels, Postgres for cases, Pinecone for vectors, Redis for cache, Kafka or Pub/Sub for events, warehouse for the forecast mart. Point at the system of record as you draw it.")
add_a("**Release path.** GitHub Actions or GitLab CI, image scan, eval gate, Helm, ArgoCD or Flux, canary, SLO check, rollback. Mark the eval gate with a pen so they see that a prompt change stops there.")

add_h("Thirty-second close")
add_a("I am an AI architect who has taken agent and data platforms into production, not only designed them. I standardize the seams, MCP, the model gateway, Kubernetes, and evaluation, and I leave teams room inside those seams. I will use Azure and AWS as landing zones, Python where the model work lives, and a typed BFF where the React client lives. I will measure grounded answers, latency, and cost, and I will not ship a tool or a corpus that fails the access-control test. The numbers I have already delivered are shorter content cycles, a lighter database, forecasts finance could use, and a delivery flow that carried less unfinished work.")


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.5)

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("AI/ML SOLUTION ARCHITECT  ·  INTERVIEW ANSWERS")
    set_run_font(hr, size=9, bold=True, color=TEAL)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_page_number(fp)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(2)
    tr = title.add_run("Interview answers")
    set_run_font(tr, size=26, bold=True, color=NAVY)

    sub = doc.add_paragraph()
    sub.paragraph_format.space_before = Pt(0)
    sub.paragraph_format.space_after = Pt(2)
    sr = sub.add_run("AI/ML Solution Architect")
    set_run_font(sr, size=16, color=TEAL)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(8)
    mr = meta.add_run("Nimmalapudi Prasad    ·    Prepared 13 October 2026    ·    45 questions")
    set_run_font(mr, size=11, color=MUTED)

    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(8)
    intro.paragraph_format.line_spacing = 1.12
    ir = intro.add_run(
        "Answers matched to the solution-architect job description: end-to-end and microservice design, "
        "enterprise integration, AI/ML and GenAI platforms, Azure and AWS, containers and CI/CD, "
        "observability, security and API management, and technical leadership. "
        "Technical examples use your production stack: LangGraph, CrewAI, MCP, FastAPI, RAG, Pinecone, "
        "Redis, Docker, Kubernetes, Helm, Terraform, ArgoCD, Prometheus, Grafana, and the model gateway "
        "across OpenAI, Claude, Gemini, and Groq."
    )
    set_run_font(ir, size=11, color=BODY)

    for item in CONTENT:
        if item[0] == "h":
            h1(doc, item[1])
        elif item[0] == "q":
            question(doc, item[1], item[2])
        elif item[0] == "a":
            label_answer(doc, item[1])

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
